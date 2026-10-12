"""Free-form inputs: PDFs, office documents, HTML, plain text, images."""

from __future__ import annotations

import io
import re
from html.parser import HTMLParser
from pathlib import Path

from .types import Converted, PageImage, TextDoc

LINES_PER_TEXT_PAGE = 60
# A PDF page with fewer extractable characters than this is treated as a scan.
MIN_PAGE_CHARS = 25
RENDER_SCALE = 2.0  # 144 dpi


def clean_line(line: str) -> str:
    """Keep indentation (it encodes hierarchy); collapse wide column gaps to save tokens."""
    line = line.replace("\t", "    ").rstrip()
    stripped = line.lstrip(" ")
    indent = min(len(line) - len(stripped), 40)
    return " " * indent + re.sub(r" {3,}", "   ", stripped)


def to_lines(text: str) -> list[str]:
    return [clean_line(ln) for ln in text.splitlines() if ln.strip()]


def paginate(lines: list[str], size: int = LINES_PER_TEXT_PAGE) -> list[list[str]]:
    return [lines[i : i + size] for i in range(0, len(lines), size)] or [[]]


def _shift_indent(lines: list[str]) -> list[str]:
    """Remove the indentation common to a whole page."""
    pad = min((len(ln) - len(ln.lstrip(" ")) for ln in lines), default=0)
    return [ln[pad:] for ln in lines] if pad else lines


COLUMN_GAP = "   "  # three spaces mark a column boundary in the text the model sees
INDENT_POINTS = 4.0  # one space of indentation per this many PDF points


def page_words(page) -> list[dict]:
    return page.extract_words(x_tolerance=1.5, y_tolerance=3, extra_attrs=["size"])


def cluster_rows(words: list[dict]) -> list[list[dict]]:
    """Group a page's words into printed lines, top to bottom."""
    rows: list[list[dict]] = []
    for word in sorted(words, key=lambda w: (round(w["top"], 1), w["x0"])):
        mid = (word["top"] + word["bottom"]) / 2
        if rows and rows[-1][-1]["top"] - 1 <= mid <= rows[-1][-1]["bottom"] + 1:
            rows[-1].append(word)
        else:
            rows.append([word])
    return rows


def line_boxes(words: list[dict]) -> list[tuple[float, float, float, float]]:
    """(x0, top, x1, bottom) in PDF points for each line `words_to_lines` returns, in order."""
    return [
        (min(w["x0"] for w in row), min(w["top"] for w in row),
         max(w["x1"] for w in row), max(w["bottom"] for w in row))
        for row in cluster_rows(words)
    ]  # fmt: skip


def words_to_lines(words: list[dict], left: float) -> list[str]:
    """Rebuild a page's lines from word positions.

    Words closer than a normal word space are joined by one space; a wider gap becomes a
    column gap. Indentation is each line's distance from the document's left margin, so
    hierarchy survives and stays comparable from page to page.
    """
    rows = cluster_rows(words)
    lines = []
    for row in rows:
        row.sort(key=lambda w: w["x0"])
        parts = [row[0]["text"]]
        for prev, word in zip(row, row[1:], strict=False):
            gap = word["x0"] - prev["x1"]
            size = float(word.get("size") or prev.get("size") or 10)
            parts.append(COLUMN_GAP if gap > 0.85 * size else " " if gap > -0.5 else "")
            parts.append(word["text"])
        indent = min(max(int((row[0]["x0"] - left) / INDENT_POINTS), 0), 40)
        lines.append(" " * indent + "".join(parts))
    return [ln for ln in lines if ln.strip()]


def convert_pdf(path: Path, file_type: str) -> Converted:
    import pdfplumber

    scans: list[int] = []
    warnings: list[str] = []
    per_page: list[list[dict]] = []
    has_images: list[bool] = []
    with pdfplumber.open(path) as pdf:
        for i, page in enumerate(pdf.pages):
            try:
                per_page.append(page_words(page))
            except Exception as exc:
                warnings.append(f"page {i + 1}: text extraction failed ({exc})")
                per_page.append([])
            has_images.append(bool(page.images))
            page.flush_cache()
    left = min((w["x0"] for words in per_page for w in words), default=0.0)
    pages = [words_to_lines(words, left) for words in per_page]
    for i, lines in enumerate(pages):
        if sum(len(ln.strip()) for ln in lines) < MIN_PAGE_CHARS and has_images[i]:
            scans.append(i)
    if not pages:
        return Converted("empty", file_type, "pdfplumber", note="PDF has no pages")
    images = render_pdf_pages(path, scans) if scans else []
    doc = TextDoc(pages, "pdf_text")
    if doc.n_lines == 0 and not images:
        return Converted("empty", file_type, "pdfplumber", note="PDF has no text and no images")
    return Converted("text", file_type, "pdfplumber", text=doc, images=images, warnings=warnings)


def render_pdf_pages(path: Path, page_indexes: list[int]) -> list[PageImage]:
    import pypdfium2 as pdfium

    out = []
    pdf = pdfium.PdfDocument(str(path))
    try:
        for i in page_indexes:
            bitmap = pdf[i].render(scale=RENDER_SCALE)
            buf = io.BytesIO()
            bitmap.to_pil().convert("RGB").save(buf, format="PNG", optimize=True)
            out.append(PageImage(page=i + 1, png=buf.getvalue()))
    finally:
        pdf.close()
    return out


def convert_image(path: Path, file_type: str) -> Converted:
    from PIL import Image

    images = []
    with Image.open(path) as img:
        for i in range(getattr(img, "n_frames", 1)):
            img.seek(i)
            frame = img.convert("RGB")
            frame.thumbnail((2200, 2200))
            buf = io.BytesIO()
            frame.save(buf, format="PNG", optimize=True)
            images.append(PageImage(page=i + 1, png=buf.getvalue()))
    doc = TextDoc([[] for _ in images], "ocr")
    return Converted("text", file_type, "image", text=doc, images=images)


def convert_plain_text(text: str, file_type: str) -> Converted:
    lines = to_lines(text)
    if not lines:
        return Converted("empty", file_type, "text", note="file is empty")
    return Converted("text", file_type, "text", text=TextDoc(paginate(lines), "text"))


def _table_lines(rows: list[list[str]]) -> list[str]:
    return [" | ".join(c.strip().replace("\n", " ") for c in row) for row in rows if any(row)]


def convert_docx(path: Path, file_type: str) -> Converted:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = docx.Document(str(path))
    lines: list[str] = []
    for block in document.iter_inner_content():
        if isinstance(block, Paragraph):
            lines += to_lines(block.text)
        elif isinstance(block, Table):
            lines += _table_lines([[c.text for c in row.cells] for row in block.rows])
    if not lines:
        return Converted("empty", file_type, "python-docx", note="document has no text")
    return Converted("text", file_type, "python-docx", text=TextDoc(paginate(lines), "office"))


def convert_pptx(path: Path, file_type: str) -> Converted:
    from pptx import Presentation
    from pptx.util import Emu

    presentation = Presentation(str(path))
    pages: list[list[str]] = []
    pictures = 0

    def walk(shapes, lines: list[str]) -> None:
        nonlocal pictures
        ordered = sorted(shapes, key=lambda s: (s.top or Emu(0), s.left or Emu(0)))
        for shape in ordered:
            if shape.shape_type == 6:  # group
                walk(shape.shapes, lines)
            elif getattr(shape, "has_table", False) and shape.has_table:
                lines.extend(_table_lines([[c.text for c in r.cells] for r in shape.table.rows]))
            elif shape.has_text_frame:
                lines.extend(to_lines(shape.text_frame.text))
            elif shape.shape_type == 13:
                pictures += 1

    for slide in presentation.slides:
        lines: list[str] = []
        walk(slide.shapes, lines)
        pages.append(lines)
    warnings = [f"{pictures} picture(s) on slides were not read"] if pictures else []
    doc = TextDoc(pages, "office")
    if doc.n_lines == 0:
        return Converted(
            "empty", file_type, "python-pptx", note="slides have no text", warnings=warnings
        )
    return Converted("text", file_type, "python-pptx", text=doc, warnings=warnings)


class _HTMLText(HTMLParser):
    _BLOCK = {"p", "div", "br", "tr", "li", "h1", "h2", "h3", "h4", "h5", "h6", "table", "section"}

    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []
        self._skip = 0

    def handle_starttag(self, tag, attrs):
        if tag in ("script", "style"):
            self._skip += 1
        elif tag in self._BLOCK:
            self.parts.append("\n")
        elif tag in ("td", "th"):
            self.parts.append(" | ")

    def handle_endtag(self, tag):
        if tag in ("script", "style"):
            self._skip = max(0, self._skip - 1)
        elif tag in self._BLOCK:
            self.parts.append("\n")

    def handle_data(self, data):
        if not self._skip:
            self.parts.append(data)


def convert_html(text: str, file_type: str) -> Converted:
    parser = _HTMLText()
    parser.feed(text)
    lines = [re.sub(r"^\s*\|\s*", "", ln) for ln in to_lines("".join(parser.parts))]
    lines = [ln for ln in lines if ln.strip()]
    if not lines:
        return Converted("empty", file_type, "html", note="page has no text")
    return Converted("text", file_type, "html", text=TextDoc(paginate(lines), "html"))
